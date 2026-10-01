import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, cast
from uuid import uuid4

import pytest
from docx import Document
from pydantic import SecretStr

from app.core.config import Settings
from app.core.errors import ApplicationError
from app.modules.knowledge.dataset import load_dataset
from app.modules.knowledge.models import (
    GraphEvidence,
    RelatedEntity,
    RelationsPage,
    SupportedCitation,
)
from app.modules.student.auth import AuthService
from app.modules.student.extraction import (
    DOCX,
    PDF,
    detect_mentions,
    extract_text,
    safe_filename,
    validate_document,
)
from app.modules.student.models import GapAnalysisRun, SkillEvidence, StudentProfile
from app.modules.student.service import StudentService
from app.modules.student.storage import LocalResumeStorage


def synthetic_pdf(text_lines: list[str] | None = None) -> bytes:
    lines = text_lines or [
        "SKILLS",
        "Python, SQL, MysteryTool",
        "PROJECTS",
        "Built REST APIs with Python",
    ]
    commands = ["BT /F1 12 Tf 72 720 Td"]
    for index, line in enumerate(lines):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(("" if index == 0 else "0 -18 Td ") + f"({escaped}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode())
    content.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(content)


def synthetic_docx() -> bytes:
    document = Document()
    document.add_heading("Technical Skills", level=1)
    document.add_paragraph("Postgres, pytest, Unknown Tool")
    document.add_heading("Projects", level=1)
    document.add_paragraph("Created automated tests with pytest")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def canonical_skills() -> list[Any]:
    return [entity for entity in load_dataset().entities if entity.kind == "Skill"]


def test_actual_pdf_extraction_and_conservative_normalization() -> None:
    content = synthetic_pdf()
    assert validate_document("synthetic.pdf", PDF, content, len(content))[1] == "pdf"
    text = extract_text(content, "pdf", 2, 10_000)
    mentions = detect_mentions(text, canonical_skills())
    mapped = {(item.raw_text, item.skill_id, item.section) for item in mentions}
    assert ("Python", "skill_python", "SKILLS") in mapped
    assert ("SQL", "skill_sql", "SKILLS") in mapped
    assert any(item.raw_text == "MysteryTool" and item.skill_id is None for item in mentions)
    assert sum(item.skill_id == "skill_python" for item in mentions) == 2


def test_actual_docx_extraction_alias_and_unknown() -> None:
    content = synthetic_docx()
    assert validate_document("synthetic.docx", DOCX, content, len(content))[1] == "docx"
    mentions = detect_mentions(extract_text(content, "docx", 20, 10_000), canonical_skills())
    postgres = next(item for item in mentions if item.raw_text == "Postgres")
    assert postgres.skill_id == "skill_postgresql"
    assert postgres.normalization_status == "ALIAS"
    assert any(item.normalization_status == "UNRESOLVED" for item in mentions)


def test_skill_section_bullets_preserve_unknown_mentions() -> None:
    mentions = detect_mentions("TECHNICAL SKILLS\nPython • MysteryTool", canonical_skills())
    assert any(item.skill_id == "skill_python" for item in mentions)
    assert any(
        item.raw_text == "MysteryTool" and item.normalization_status == "UNRESOLVED"
        for item in mentions
    )


@pytest.mark.parametrize(
    "filename,media,content,maximum,error",
    [
        ("../private.pdf", PDF, synthetic_pdf(), 99999, "Unsafe filename"),
        ("resume.exe", PDF, synthetic_pdf(), 99999, "UNSUPPORTED"),
        ("resume.pdf", DOCX, synthetic_pdf(), 99999, "UNSUPPORTED"),
        ("resume.pdf", PDF, b"", 99999, "EMPTY"),
        ("resume.pdf", PDF, synthetic_pdf(), 10, "TOO_LARGE"),
        ("resume.pdf", PDF, b"%PDF-broken", 99999, None),
        ("resume.docx", DOCX, b"PK broken", 99999, "CORRUPT"),
    ],
)
def test_upload_validation_boundaries(
    filename: str, media: str, content: bytes, maximum: int, error: str | None
) -> None:
    if error is None:
        validate_document(filename, media, content, maximum)
    else:
        with pytest.raises(ValueError, match=error):
            validate_document(filename, media, content, maximum)


async def test_private_storage_rejects_paths_and_round_trips() -> None:
    storage = LocalResumeStorage(Path(__file__).parents[2] / "artifacts/phase3/unit-storage-v2")
    await storage.initialize()
    with pytest.raises(ValueError):
        storage.path("../escape.pdf")
    key = f"resume_{uuid4().hex}.pdf"
    await storage.save(key, b"private")
    assert await storage.read(key) == b"private"
    await storage.delete(key)
    assert not await storage.exists(key)


class AuthRepository:
    async def account_by_email(self, email: str) -> dict[str, object] | None:
        return None


def test_jwt_identity_is_signed_and_rejects_tampering() -> None:
    settings = Settings(_env_file=None, jwt_secret=SecretStr("x" * 32))
    auth = AuthService(settings, cast(Any, AuthRepository()))
    student = f"student_{'1' * 32}"
    token = auth.token(student)
    assert auth.identity(token.access_token).student_id == student
    with pytest.raises(ApplicationError) as raised:
        auth.identity(token.access_token + "broken")
    assert raised.value.code == "AUTHENTICATION_REQUIRED"


class GapRepository:
    def __init__(self, evidence: list[SkillEvidence]) -> None:
        self.values = evidence
        self.saved: GapAnalysisRun | None = None

    async def profile(self, student_id: str) -> StudentProfile:
        now = datetime.now(UTC)
        return StudentProfile(
            student_id=student_id,
            display_name="Synthetic",
            target_role_id="role_backend_developer",
            created_at=now,
            updated_at=now,
            version=3,
        )

    async def evidence(self, student_id: str, resume_id: str | None = None) -> list[SkillEvidence]:
        return self.values

    async def save_gap(self, run: GapAnalysisRun) -> None:
        self.saved = run


class GapKnowledge:
    async def related(self, *args: object, **kwargs: object) -> RelationsPage:
        dataset = load_dataset()
        entities = {entity.id: entity for entity in dataset.entities}
        sources = {source.id: source for source in dataset.sources}
        selected = [
            assertion
            for assertion in dataset.assertions
            if assertion.source_id == "role_backend_developer"
            and assertion.target_id in {"skill_python", "skill_sql", "skill_http"}
        ]
        return RelationsPage(
            entity=entities["role_backend_developer"],
            limit=100,
            offset=0,
            items=[
                RelatedEntity(
                    entity=entities[assertion.target_id],
                    evidence=GraphEvidence(
                        assertion=assertion,
                        dataset_version=dataset.version,
                        path=[assertion.source_id, assertion.type, assertion.target_id],
                        provenance=[
                            SupportedCitation(citation=citation, source=sources[citation.source_id])
                            for citation in assertion.provenance
                        ],
                    ),
                )
                for assertion in selected
            ],
        )


def evidence(skill: str, status: str, suffix: str) -> SkillEvidence:
    now = datetime.now(UTC)
    return SkillEvidence(
        evidence_id=f"evidence_{suffix * 32}",
        student_id=f"student_{'1' * 32}",
        resume_id=f"resume_{'2' * 32}",
        raw_text=skill,
        section="SKILLS",
        evidence_text=skill,
        skill_id=f"skill_{skill.casefold()}",
        skill_name=skill,
        normalization_status="EXACT",
        verification_status=status,
        observed_at=now,
        created_at=now,
        updated_at=now,
    )


async def test_gap_rules_supported_partial_and_unverified() -> None:
    repository = GapRepository(
        [evidence("Python", "CONFIRMED", "a"), evidence("SQL", "EXTRACTED", "b")]
    )
    service = StudentService(
        Settings(_env_file=None),
        cast(Any, repository),
        cast(Any, GapKnowledge()),
        LocalResumeStorage(Path(__file__).parents[2] / "artifacts/phase3/gap-storage"),
    )
    run = await service.analyze(f"student_{'1' * 32}")
    assert {item.skill_id: item.status for item in run.items} == {
        "skill_python": "SUPPORTED",
        "skill_sql": "PARTIALLY_SUPPORTED",
        "skill_http": "UNVERIFIED",
    }
    assert run.rule_version == "gap-rules-v1"
    assert run.knowledge_dataset_version == "careerpilot-knowledge-v1"
    assert repository.saved == run


def test_safe_filename_preserves_metadata_without_becoming_a_path() -> None:
    assert safe_filename("My Résumé 2026.pdf") == "My R_sum_ 2026.pdf"
